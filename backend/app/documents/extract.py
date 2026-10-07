"""Pull plain text out of uploaded files.

Parsers are third-party code run on untrusted input, so every failure becomes an
ExtractionError with a short, user-presentable message; internals are never shown.
"""

import io
import zipfile
from pathlib import PurePath

from docx import Document as DocxDocument
from pypdf import PdfReader

MAX_PDF_PAGES = 500
# Cap on text kept per document, so one huge file can't flood chunking/embedding
MAX_TEXT_CHARS = 2_000_000
MAX_UNPACKED_DOCX_BYTES = 50 * 1024 * 1024

# extension -> canonical content type. The client-sent Content-Type is ignored:
# it is attacker-controlled.
SUPPORTED_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ExtractionError(Exception):
    """The file can't be turned into text. `str(exc)` is safe to show the user."""


class UnsupportedFileTypeError(ExtractionError):
    pass


def file_extension(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def detect_type(filename: str, data: bytes) -> str:
    """The extension decides the type; the content must agree with it."""
    ext = file_extension(filename)
    if ext not in SUPPORTED_TYPES:
        raise UnsupportedFileTypeError(
            "Unsupported file type. Allowed: " + ", ".join(sorted(SUPPORTED_TYPES)) + "."
        )
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise UnsupportedFileTypeError("The file is not a valid PDF.")
    if ext == ".docx" and not data.startswith(b"PK"):
        raise UnsupportedFileTypeError("The file is not a valid DOCX document.")
    if ext in (".txt", ".md") and b"\x00" in data[:8192]:
        raise UnsupportedFileTypeError("The file does not look like a text file.")
    return SUPPORTED_TYPES[ext]


def extract_text(filename: str, data: bytes) -> str:
    """Blocking (CPU-bound parsing): call it through asyncio.to_thread."""
    ext = file_extension(filename)
    try:
        if ext in (".txt", ".md"):
            text = _extract_plain(data)
        elif ext == ".pdf":
            text = _extract_pdf(data)
        elif ext == ".docx":
            text = _extract_docx(data)
        else:
            raise UnsupportedFileTypeError("Unsupported file type.")
    except ExtractionError:
        raise
    except Exception as exc:  # parsers raise many exception types on bad input
        raise ExtractionError("The file could not be read; it may be corrupt.") from exc

    text = text.replace("\x00", "").strip()
    if not text:
        raise ExtractionError("No text could be extracted (is the file empty or scanned?).")
    return text[:MAX_TEXT_CHARS]


def _extract_plain(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ExtractionError("Password-protected PDFs are not supported.")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ExtractionError(f"PDFs are limited to {MAX_PDF_PAGES} pages.")
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    # Reject zip bombs before python-docx inflates the archive into memory
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED_DOCX_BYTES:
            raise ExtractionError("The document is too large when unpacked.")
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n\n".join(parts)
